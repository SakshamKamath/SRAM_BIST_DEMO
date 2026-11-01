# Two click-driven slides for the verification section, same style as isol_pptx.py / vad_obs_pptx.py.
#   slide 1: how the additions were verified (three levels, three principles)
#   slide 2: one row per feature: how it is tested, what it shows, result
# Results: Questa regression of 31 Oct (boundary scan, isolation, MBIST, Croc tests) and
# Verilator 5.020 runs of 1 Nov (VAD testbench, test_vad_irq, test_interrupts). No addresses.
# usage: python verification_pptx.py out.pptx
import sys
from lxml import etree
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

TEAL, PURPLE, AMBER, DARK, RED, GREEN = "00897B", "93005F", "E39B2D", "12343B", "C62828", "2E7D32"
GREY, LIGHT, BLUE, PALE = "9E9E9E", "E0F2F1", "1E5A8C", "F5F5F5"
SANS, MONO = "Calibri", "Consolas"
FADE = 250

prs = Presentation()
prs.slide_width, prs.slide_height = Emu(9144000), Emu(5143500)   # 10" x 5.625"
px = lambda v: Emu(int(v * 9144000 / 1280))                       # design grid: 1280 x 720
fs = lambda v: Pt(v * 0.5625)
slide = None

def box(name, x, y, w, h, fill=None, line=None, lw=1.5, lines=(), shape=MSO_SHAPE.RECTANGLE,
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, margin=4, dash=False):
    s = slide.shapes.add_shape(shape, px(x), px(y), px(w), px(h))
    s.name = name
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = min(0.15, 10 / min(w, h))
    st = s._element.find(qn("p:style"))
    if st is not None: s._element.remove(st)   # no theme shadow/effects
    if fill: s.fill.solid(); s.fill.fore_color.rgb = RGBColor.from_string(fill)
    else: s.fill.background()
    if line:
        s.line.color.rgb = RGBColor.from_string(line); s.line.width = Pt(lw)
        if dash: s.line.dash_style = 4   # dash
    else: s.line.fill.background()
    tf = s.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", px(margin))
    for k, ln in enumerate(lines):
        text, size, color = ln[:3]
        bold = ln[3] if len(ln) > 3 else False
        font = ln[4] if len(ln) > 4 else SANS
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run(); r.text = text
        r.font.size, r.font.bold, r.font.name = fs(size), bold, font
        r.font.color.rgb = RGBColor.from_string(color)
    return s

def wire(name, x1, y1, x2, y2, color=DARK, lw=2.5, arrow=True, dash=False):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, px(x1), px(y1), px(x2), px(y2))
    c.name = name
    st = c._element.find(qn("p:style"))
    if st is not None: c._element.remove(st)
    c.line.color.rgb = RGBColor.from_string(color); c.line.width = Pt(lw)
    if dash: c.line.dash_style = 4
    if arrow:
        ln = c.line._get_or_add_ln()
        etree.SubElement(ln, qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"})
    return c

def title(text):
    box("Title", 40, 14, 1200, 50, lines=[(text, 32, DARK, True)], align=PP_ALIGN.LEFT, margin=0)

def caption(name, text):
    return box(name, 0, 628, 1280, 92, fill=DARK, lines=[(s, 22, "FFFFFF") for s in text.split("\n")],
               align=PP_ALIGN.LEFT, margin=40)

def ev(c, t, shp, kind="in"):
    c.append((t, shp, kind))

def attach_timing(sl, clicks):
    ids = [2]
    def nid():
        ids[0] += 1; return ids[0]
    grp = {}                                  # (spid, kind) -> grpId
    def effect(spid, kind, node):
        g = grp.setdefault((spid, kind), 0 if kind == "in" else 1)
        eid, a, b = nid(), nid(), nid()
        tgt = f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>'
        vis = lambda cid, d, v: (f'<p:set><p:cBhvr><p:cTn id="{cid}" dur="1" fill="hold"><p:stCondLst><p:cond delay="{d}"/></p:stCondLst></p:cTn>'
                                 f'{tgt}<p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr>'
                                 f'<p:to><p:strVal val="{v}"/></p:to></p:set>')
        fade = lambda cid, d: (f'<p:animEffect transition="{d}" filter="fade"><p:cBhvr><p:cTn id="{cid}" dur="{FADE}"/>{tgt}</p:cBhvr></p:animEffect>')
        body = vis(a, 0, "visible") + fade(b, "in") if kind == "in" else fade(a, "out") + vis(b, FADE - 1, "hidden")
        cls = "entr" if kind == "in" else "exit"
        return (f'<p:par><p:cTn id="{eid}" presetID="10" presetClass="{cls}" presetSubtype="0" fill="hold" grpId="{g}" nodeType="{node}">'
                f'<p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>{body}</p:childTnLst></p:cTn></p:par>')

    seq = []
    for c in clicks:
        times = sorted({t for t, _, _ in c})
        outer = nid()
        inner = []
        for gi, t in enumerate(times):
            pid = nid()
            effs = [effect(s.shape_id, k, ("clickEffect" if gi == 0 else "afterEffect") if j == 0 else "withEffect")
                    for j, (_, s, k) in enumerate([x for x in c if x[0] == t])]
            inner.append(f'<p:par><p:cTn id="{pid}" fill="hold"><p:stCondLst><p:cond delay="{t}"/></p:stCondLst>'
                         f'<p:childTnLst>{"".join(effs)}</p:childTnLst></p:cTn></p:par>')
        seq.append(f'<p:par><p:cTn id="{outer}" fill="hold"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst>'
                   f'<p:childTnLst>{"".join(inner)}</p:childTnLst></p:cTn></p:par>')
    bld = "".join(f'<p:bldP spid="{s}" grpId="{g}" animBg="1"/>' for (s, _), g in grp.items())
    NS = 'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"'
    timing = (f'<p:timing {NS}><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>'
              '<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>'
              + "".join(seq) +
              '</p:childTnLst></p:cTn><p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
              '<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>'
              f'</p:childTnLst></p:cTn></p:par></p:tnLst><p:bldLst>{bld}</p:bldLst></p:timing>')
    sl._element.append(etree.fromstring(timing))

def chip(name, x, y, text, color):
    return box(name, x, y, 230, 28, fill=color, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
               lines=[(text, 15, "FFFFFF", True)], margin=0)

# =====================================================================================

# =====================================================================================
# Slide 1: approach
# =====================================================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
title("How the additions were verified")
box("Subtitle", 40, 66, 1200, 34, lines=[(
    "Every feature is checked on its own, then with software on the CPU, then from the chip pins like a tester.",
    18, GREY)], align=PP_ALIGN.LEFT, margin=0)

LV = [("1", "Unit", "One block with simple models around it", "VAD registers, March BIST with injected memory faults", TEAL),
      ("2", "SoC", "Real C programs running on the CPU", "Memory BIST, VAD interrupt, original Croc tests", BLUE),
      ("3", "Chip pins", "Only the pins and JTAG, as on a tester", "Boundary scan, crossbar isolation, sub access", PURPLE)]
clicks = [[] for _ in range(4)]
LW, LY, LH = 360, 124, 250
for k, (num, head, how, ex, col) in enumerate(LV):
    x = 40 + k * (LW + 60)
    c = clicks[k]
    for s in (box(f"Level {num}", x, LY, LW, LH, fill="FFFFFF", line=col, lw=2.5, shape=MSO_SHAPE.ROUNDED_RECTANGLE),
              box(f"Level {num} num", x + 18, LY + 18, 44, 44, fill=col, shape=MSO_SHAPE.OVAL,
                  lines=[(num, 22, "FFFFFF", True)], margin=0),
              box(f"Level {num} head", x + 76, LY + 18, LW - 90, 44, lines=[(head, 26, DARK, True)],
                  align=PP_ALIGN.LEFT, margin=0),
              box(f"Level {num} how", x + 18, LY + 76, LW - 36, 60, lines=[(how, 20, DARK)],
                  align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, margin=0),
              box(f"Level {num} ex", x + 18, LY + 150, LW - 36, 90, lines=[("Used for", 15, col, True), (ex, 17, GREY)],
                  align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, margin=0)):
        ev(c, 0, s)
    if k:
        ev(c, 0, wire(f"Arrow {k}", x - 54, LY + LH / 2, x - 6, LY + LH / 2, DARK, 3))

PR = [("Self-checking", "Every check prints pass or fail, and each run ends with a count. Nothing is judged by eye."),
      ("Used as in silicon", "Each feature is driven the way it will be used: by software over the bus, by a tester over JTAG, or both."),
      ("Negative controls", "Deliberately broken hardware or wrong expectations must make the test fail, and they do.")]
PW, PY, PH = 380, 420, 180
c = clicks[3]
for k, (head, txt) in enumerate(PR):
    x = 40 + k * (PW + 30)
    for s in (box(f"Principle {k}", x, PY, PW, PH, fill=LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE),
              box(f"Principle {k} head", x + 20, PY + 18, PW - 40, 36, lines=[(head, 22, TEAL, True)],
                  align=PP_ALIGN.LEFT, margin=0),
              box(f"Principle {k} text", x + 20, PY + 62, PW - 40, PH - 76, lines=[(txt, 19, DARK)],
                  align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, margin=0)):
        ev(c, 0, s)
attach_timing(slide, clicks)
slide.notes_slide.notes_text_frame.text = (
    "Clicks 1-3: the three levels. Unit: a testbench around one block (tb_vad_sram_dbg, tb_march_bist_top with a "
    "behavioural SRAM fault model). SoC: C programs on the CVE2 core in tb_croc_soc (test_mbist, test_vad_mbist, "
    "test_vad_irq, the original Croc tests). Chip: tb_croc_chip, only the pads; JTAG drives boundary scan "
    "(+bscan_test) and crossbar isolation / sub access (+obi_jtag_test). Click 4: the principles. Negative controls: "
    "the VAD testbench was run against the RTL without the new interrupt bits and failed at those checks; "
    "test_vad_irq with one wrong expectation reports FAIL with that check's number.")

# =====================================================================================
# Slide 2: per feature
# =====================================================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
title("Verification per feature")
box("Subtitle", 40, 66, 1200, 34, lines=[(
    "Self-checking runs in RTL simulation. Counts are checks passed / checks run.", 18, GREY)],
    align=PP_ALIGN.LEFT, margin=0)

CX = [(40, 230, "Feature"), (280, 170, "Tested at"), (460, 620, "What the test shows"), (1090, 150, "Result")]
HY, RY, RH = 108, 146, 54
for x, w, t in CX:
    box(f"Head {t}", x, HY, w, 32, lines=[(t, 16, GREY, True)], align=PP_ALIGN.LEFT, margin=0)
wire("Head rule", 40, HY + 34, 1240, HY + 34, GREY, 1, arrow=False)

OK, TODO = GREEN, AMBER
ROWS = [
    ("Boundary scan", "Chip pins", "Reads and drives every pin, feeds the core from the scan cells, then hands the pins back", "34 / 34", OK),
    ("March BIST", "Unit, SoC, chip", "All memories pass from software and JTAG; unit test runs with injected memory faults", "SoC + chip pass", OK),
    ("Crossbar isolation", "Chip pins", "Each route and the error path, a remapped address map, nothing leaks while isolated", "18 / 18", OK),
    ("Subordinate access", "Chip pins", "JTAG writes and reads a real memory; each transfer arrives exactly once; errors come back", "25 / 25", OK),
    ("VAD observability", "Unit", "Frame memory access; energy and zero crossings match a model; decision on chosen values", "1,095 / 1,095", OK),
    ("VAD interrupt", "Unit, SoC", "Fired on demand, seen in the VAD and the CPU, taken once, cleared by software", "Pass", OK),
    ("Testability", "ATPG", "Same netlist, registers reachable vs. cut off: pin-level fault coverage 17 % → 28 %", "+11 points", OK),
    ("No regressions", "SoC, chip", "The original Croc tests still pass with every addition in place", "10 / 10", OK),
]
GROUPS = [range(0, 4), range(4, 6), range(6, 8)]
clicks = [[] for _ in range(len(GROUPS) + 1)]
for g, idx in enumerate(GROUPS):
    for i in idx:
        feat, lvl, shows, res, rc = ROWS[i]
        y = RY + i * RH
        c = clicks[g]
        shapes = [box(f"Row {i} bg", 40, y + 3, 1200, RH - 6, fill=PALE if i % 2 == 0 else "FFFFFF",
                      shape=MSO_SHAPE.ROUNDED_RECTANGLE),
                  box(f"Row {i} feat", 52, y, 220, RH, lines=[(feat, 18, DARK, True)], align=PP_ALIGN.LEFT, margin=0),
                  box(f"Row {i} lvl", 280, y, 170, RH, lines=[(lvl, 16, GREY)], align=PP_ALIGN.LEFT, margin=0),
                  box(f"Row {i} shows", 460, y, 620, RH, lines=[(shows, 16, DARK)], align=PP_ALIGN.LEFT, margin=0),
                  box(f"Row {i} res", 1090, y + 12, 140, RH - 24, fill=rc, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
                      lines=[(res, 15, "FFFFFF", True)], margin=0)]
        for s in shapes: ev(c, 0, s)
ev(clicks[-1], 0, caption("Limits",
    "RTL simulation only: no gate-level or tester runs yet. The boundary-scan and isolation runs are older than the\n"
    "last RTL edits and need a rerun; the BIST fault-injection unit results are still to be added."))
attach_timing(slide, clicks)
slide.notes_slide.notes_text_frame.text = (
    "Click 1: the DFT features. Boundary scan: tb_croc_chip +bscan_test, 34 checks (chain lengths for BYPASS / "
    "SAMPLE / EXTEST / INTEST, walking patterns, INTEST with the core echoing inputs, release). March BIST: "
    "test_mbist and test_vad_mbist pass on tb_croc_soc and tb_croc_chip; the unit test tb_march_bist_top injects "
    "stuck-at, transition and read-disturb faults (result to be filled in). Crossbar isolation: tb_croc_chip "
    "+obi_jtag_test, 18 checks; with subordinate access 25 checks including the kept JTAG map. "
    "Click 2: VAD. tb_vad_sram_dbg, 1095 checks including the interrupt trigger and read-back; test_vad_irq on the "
    "SoC checks the VAD status bit, the CPU pending bit and the handler. Click 3: ATPG ablation on the DEMO netlist, "
    "28.3 % vs 17.0 % pins-only fault coverage. Regressions: helloworld, print_config, clint, compute, gpio, "
    "obi_timer, soc_ctrl, uart, interrupts on the SoC and helloworld on the chip; test_idma not re-run. "
    "Simulators: Questa on 31 Oct; Verilator 5.020 on 1 Nov for the VAD testbench, test_vad_irq and test_interrupts.")

prs.save(sys.argv[1])
print("saved", sys.argv[1])
